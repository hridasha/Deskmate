"""Retrieval-augmented generation over the user's own files.

Pipeline: extract text (with page/sheet locators) -> chunk with overlap -> embed -> store.
Search: semantic similarity and BM25 keyword scoring, merged with reciprocal-rank fusion, so exact terms
(names, numbers, part codes) and paraphrases both work. Results carry the source name and location for citation.
"""
import asyncio
import base64
import csv
import hashlib
import html
import io
import json
import math
import re
import time
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import chunking
import db
import embeddings
import llm
import memory
import rerank
import secrets_store
from router import tokens

MAX_BYTES = 15 * 1024 * 1024
FILES_DIR = secrets_store.DATA_DIR / "files"  # the original uploads, so an answer can link back to the real file
MAX_DOCS = 25
MAX_CHUNKS_PER_DOC = 600
MAX_CHUNKS_PER_USER = 4000
CHUNK_CHARS, OVERLAP = 900, 120
MAX_PDF_PAGES, MAX_ROWS = 500, 20000

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".json", ".log", ".yaml", ".yml", ".xml", ".sql", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".go", ".rs", ".ini", ".toml"}
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp"}
SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".html", ".htm"} | TEXT_EXT | set(IMAGE_MIME)
SUPPORTED_LABEL = "PDF (including scans), images (PNG/JPG, read with AI), Word (.docx), Excel (.xlsx), CSV/TSV, HTML, and text or code files (.txt, .md, .json, ...)"
MAX_OCR_BYTES = 8 * 1024 * 1024


class IngestError(ValueError):
    """A problem with the uploaded file that the user can act on."""


class NeedsOCR(Exception):
    """The PDF has no text layer (a scan), so it must be read as images."""


@dataclass
class Segment:
    text: str
    locator: str


# ---------------- extraction ----------------


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _rows_to_segments(rows, label: str) -> list[Segment]:
    """Group table rows into segments that repeat the header, so each chunk is self-explanatory."""
    out: list[Segment] = []
    header: list[str] | None = None
    batch: list[str] = []
    start, n = 1, 0
    for row in rows:
        cells = ["" if c is None else str(c).strip() for c in row]
        if not any(cells):
            continue
        n += 1
        if n > MAX_ROWS:
            break
        if header is None:
            header = [c or f"col{i + 1}" for i, c in enumerate(cells)]
            continue
        batch.append("; ".join(f"{h}: {c}" for h, c in zip(header, cells) if c))
        if len(batch) == 25:
            out.append(Segment("\n".join(batch), f"{label} rows {start}-{start + len(batch) - 1}"))
            start += len(batch)
            batch = []
    if batch:
        out.append(Segment("\n".join(batch), f"{label} rows {start}-{start + len(batch) - 1}"))
    if header and not out:
        out.append(Segment("Columns: " + ", ".join(header), label))
    return out


def extract_pdf(data: bytes) -> list[Segment]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestError("This PDF is password-protected.")
        segs = []
        for i, page in enumerate(reader.pages[:MAX_PDF_PAGES]):
            text = (page.extract_text() or "").strip()
            if text:
                segs.append(Segment(text, f"p.{i + 1}"))
    except PdfReadError as e:
        raise IngestError("Couldn't read this PDF; it may be damaged.") from e
    if not segs:
        raise NeedsOCR()
    return segs


def extract_docx(data: bytes) -> list[Segment]:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 - python-docx raises various errors for bad files
        raise IngestError("Couldn't read this Word file; it may be damaged.") from e
    segs: list[Segment] = []
    heading, buf = "start", []

    def flush():
        if buf:
            segs.append(Segment("\n".join(buf), f"section '{heading}'"))
            buf.clear()

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if (p.style.name or "").lower().startswith("heading"):
            flush()
            heading = text[:60]
        buf.append(text)
    flush()
    for ti, table in enumerate(doc.tables):
        lines = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if any(lines):
            segs.append(Segment("\n".join(lines), f"table {ti + 1}"))
    return segs


def extract_xlsx(data: bytes) -> list[Segment]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise IngestError("Couldn't read this spreadsheet; it may be damaged or password-protected.") from e
    segs: list[Segment] = []
    for ws in wb.worksheets:
        segs.extend(_rows_to_segments(ws.iter_rows(values_only=True), f"sheet '{ws.title}'"))
    return segs


def extract_csv(data: bytes, delimiter: str) -> list[Segment]:
    return _rows_to_segments(csv.reader(io.StringIO(_decode(data)), delimiter=delimiter), "rows")


def extract_html(data: bytes) -> list[Segment]:
    text = _decode(data)
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return [Segment(html.unescape(re.sub(r"\s+", " ", text)).strip(), "page")]


def extract(ext: str, data: bytes) -> list[Segment]:
    if ext == ".pdf":
        segs = extract_pdf(data)
    elif ext == ".docx":
        segs = extract_docx(data)
    elif ext in (".xlsx", ".xlsm"):
        segs = extract_xlsx(data)
    elif ext == ".csv":
        segs = extract_csv(data, ",")
    elif ext == ".tsv":
        segs = extract_csv(data, "\t")
    elif ext in (".html", ".htm"):
        segs = extract_html(data)
    else:
        segs = [Segment(_decode(data), "text")]
    return [s for s in segs if s.text.strip()]


# ---------------- OCR (scans and images) ----------------

OCR_PROMPT = (
    "Transcribe all the text in this file exactly as written, in reading order. Put tables in rows with columns separated by ' | '. "
    "For a multi-page document, start each page with a line '=== Page N ==='. If the file is an image without text (a chart or diagram), "
    "give a one-paragraph description instead, starting with 'Description:'. If there is nothing readable at all, reply exactly: NO_TEXT. "
    "Do not add commentary."
)


async def _ocr_call(data: bytes, mime: str) -> str:
    """Ask a vision-capable model to transcribe the file. Separate so tests can replace it."""
    from langchain_core.messages import HumanMessage

    msg = HumanMessage(content=[{"type": "text", "text": OCR_PROMPT}, {"type": "media", "mime_type": mime, "data": base64.b64encode(data).decode()}])
    resp = await llm.ainvoke_retry(llm.get_llm(tier="smart"), [msg])
    c = resp.content
    return c if isinstance(c, str) else "".join(p.get("text", "") if isinstance(p, dict) else str(p) for p in c)


async def ocr_segments(data: bytes, mime: str) -> list[Segment]:
    if len(data) > MAX_OCR_BYTES:
        raise IngestError(f"Scanned files and images over {MAX_OCR_BYTES // (1024 * 1024)} MB can't be read yet. Try a smaller or lower-resolution file.")
    if not llm.keys_for("gemini"):
        raise IngestError("Reading scans and images needs a Gemini API key. Add one in Settings.")
    try:
        text = (await _ocr_call(data, mime)).strip()
    except Exception as e:  # noqa: BLE001
        raise IngestError(f"Couldn't read this scan right now ({type(e).__name__}). Try again in a moment.") from e
    if not text or text == "NO_TEXT":
        raise IngestError("No readable text found in this scan.")
    parts = re.split(r"^=== Page (\d+) ===\s*$", text, flags=re.M)
    segs = [Segment(parts[0].strip(), "scan")] if parts[0].strip() else []
    for i in range(1, len(parts), 2):
        if parts[i + 1].strip():
            segs.append(Segment(parts[i + 1].strip(), f"p.{parts[i]} (scanned)"))
    return segs or [Segment(text, "scan")]


# ---------------- chunking ----------------


def chunk(segments: list[Segment]) -> list[tuple[str, str]]:
    """Split segments into ~900-character chunks that overlap, preferring to break at paragraph or sentence ends."""
    out: list[tuple[str, str]] = []
    for seg in segments:
        text = re.sub(r"[ \t]+", " ", seg.text).strip()
        if len(text) <= CHUNK_CHARS:
            out.append((text, seg.locator))
            continue
        start, first = 0, True
        while start < len(text):
            end = min(len(text), start + CHUNK_CHARS)
            if end < len(text):
                window = text[start:end]
                cut = max(window.rfind("\n"), window.rfind(". "))
                if cut > CHUNK_CHARS * 0.5:
                    end = start + cut + 1
            out.append((text[start:end].strip(), seg.locator if first else f"{seg.locator} (cont.)"))
            first = False
            if end >= len(text):
                break
            start = max(end - OVERLAP, start + 1)
    return [c for c in out if len(c[0]) > 20]


# ---------------- embeddings ----------------


async def embed_many(texts: list[str]) -> list[list[float] | None]:
    """Embed document chunks in batches. If embedding fails the chunk is still stored and found by keywords."""
    return await embeddings.embed_documents(texts)


# ---------------- storage ----------------

_cache: dict[str, tuple[int, list[dict]]] = {}
_version: Counter = Counter()


def _touch(user: str) -> None:
    _version[user] += 1


def _file_path(doc_id: str, ext: str) -> Path:
    return FILES_DIR / f"{doc_id}{ext}"


async def ingest(user: str, filename: str, data: bytes, scope: list[str] | None = None, cfg: chunking.Config | None = None, origin: dict | None = None) -> dict:
    name = Path(filename or "upload").name[:120]
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise IngestError(f"'{ext or 'this file type'}' isn't supported. Supported: {SUPPORTED_LABEL}.")
    if not data:
        raise IngestError("This file is empty.")
    if len(data) > MAX_BYTES:
        raise IngestError(f"This file is larger than {MAX_BYTES // (1024 * 1024)} MB.")

    digest = hashlib.sha256(data).hexdigest()
    existing = db.q("SELECT id, hash, scope FROM kdocs WHERE user=? AND name=?", (user, name))  # the same name is the same document, even if its contents changed
    if existing and existing[0]["hash"] == digest and not (cfg and cfg.strategy != "auto"):
        db.run("UPDATE kdocs SET synced=? WHERE id=?", (time.time(), existing[0]["id"]))
        doc = get_doc(user, existing[0]["id"]) or {}
        return {**doc, "unchanged": True}
    if existing and scope is None and existing[0]["scope"]:
        scope = json.loads(existing[0]["scope"])  # keep who may see it when a document is updated
    if not existing and db.q("SELECT COUNT(*) AS n FROM kdocs WHERE user=?", (user,))[0]["n"] >= MAX_DOCS:
        raise IngestError(f"You can keep up to {MAX_DOCS} documents. Delete one first.")

    if ext in IMAGE_MIME:
        segments = await ocr_segments(data, IMAGE_MIME[ext])
    else:
        try:
            segments = extract(ext, data)
        except NeedsOCR:  # a scanned PDF: read it as images
            segments = await ocr_segments(data, "application/pdf")
    try:
        chunks, chunk_info = await chunking.build_chunks(segments, (cfg or chunking.Config()), ext.lstrip("."))
    except ValueError as e:
        raise IngestError(str(e)) from e
    chunks = chunks[:MAX_CHUNKS_PER_DOC]
    if not chunks:
        raise IngestError("No readable text found in this file.")
    used = db.q("SELECT COUNT(*) AS n FROM kchunks WHERE doc_id IN (SELECT id FROM kdocs WHERE user=?)", (user,))[0]["n"]
    if used + len(chunks) > MAX_CHUNKS_PER_USER:
        raise IngestError("Your knowledge base is full. Delete a document to make room.")

    embs = await embed_many([c[0] for c in chunks])
    for old in existing:  # re-uploading the same file replaces it
        delete_doc(user, old["id"])
    doc_id = uuid.uuid4().hex[:10]
    model = embeddings.model_id()
    db.run("INSERT INTO kdocs (id, user, name, kind, size, chunks, embedded, scope, created, chunking, hash, origin, synced) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", (doc_id, user, name, ext.lstrip("."), len(data), len(chunks), sum(1 for e in embs if e), json.dumps(scope) if scope else None, time.time(), json.dumps(chunk_info), digest, json.dumps(origin or {"type": "upload"}), time.time()))
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    _file_path(doc_id, ext).write_bytes(data)
    _save_segments(doc_id, segments)
    with db._lock:
        conn = db.get()
        conn.executemany("INSERT INTO kchunks (doc_id, ord, text, locator, emb, emb_model) VALUES (?,?,?,?,?,?)", [(doc_id, i, c[0], c[1], json.dumps(e) if e else None, model if e else None) for i, (c, e) in enumerate(zip(chunks, embs))])
        conn.commit()
    _touch(user)
    return get_doc(user, doc_id) or {}


def get_doc(user: str, doc_id: str) -> dict | None:
    rows = db.q("SELECT * FROM kdocs WHERE user=? AND id=?", (user, doc_id))
    return _public(rows[0]) if rows else None


def _public(row: dict) -> dict:
    return {
        **{k: row[k] for k in ("id", "name", "kind", "size", "chunks", "embedded", "created")},
        "origin": json.loads(row["origin"]) if row.get("origin") else {"type": "upload"},
        "synced": row.get("synced"),
        "scope": json.loads(row["scope"]) if row["scope"] else None,
        "chunking": json.loads(row["chunking"]) if row.get("chunking") else None,
    }


def list_docs(user: str) -> list[dict]:
    return [_public(r) for r in db.q("SELECT * FROM kdocs WHERE user=? ORDER BY created DESC", (user,))]


def delete_doc(user: str, doc_id: str) -> bool:
    if not db.q("SELECT 1 FROM kdocs WHERE user=? AND id=?", (user, doc_id)):
        return False
    row = db.q("SELECT name FROM kdocs WHERE id=?", (doc_id,))
    db.run("DELETE FROM kchunks WHERE doc_id=?", (doc_id,))
    db.run("DELETE FROM ksegments WHERE doc_id=?", (doc_id,))
    db.run("DELETE FROM kdocs WHERE id=?", (doc_id,))
    if row:
        for f in FILES_DIR.glob(f"{doc_id}.*"):
            f.unlink(missing_ok=True)
    _touch(user)
    return True


def original_file(user: str, doc_id: str) -> tuple[Path, str] | None:
    """The uploaded file and its name, for 'open the source' links."""
    rows = db.q("SELECT name FROM kdocs WHERE user=? AND id=?", (user, doc_id))
    if not rows:
        return None
    found = next(iter(FILES_DIR.glob(f"{doc_id}.*")), None)
    return (found, rows[0]["name"]) if found else None


def set_scope(user: str, doc_id: str, agents: list[str] | None) -> bool:
    if not db.q("SELECT 1 FROM kdocs WHERE user=? AND id=?", (user, doc_id)):
        return False
    db.run("UPDATE kdocs SET scope=? WHERE id=?", (json.dumps(agents) if agents else None, doc_id))
    _touch(user)
    return True


def doc_names(user: str, agent_id: str, limit: int = 8, only: set[str] | None = None) -> list[str]:
    """Names of the documents this agent can search, so it knows when to look."""
    names = []
    for d in list_docs(user):
        if only is not None and d["id"] not in only:
            continue
        if d["scope"] is None or agent_id in d["scope"] or only is not None:
            names.append(d["name"])
    return names[:limit]


# ---------------- search ----------------


def _load(user: str) -> list[dict]:
    hit = _cache.get(user)
    if hit and hit[0] == _version[user]:
        return hit[1]
    rows = db.q("SELECT c.id, c.doc_id, c.text, c.locator, c.emb, c.emb_model, d.name, d.scope FROM kchunks c JOIN kdocs d ON d.id = c.doc_id WHERE d.user=?", (user,))
    current = embeddings.model_id()
    for r in rows:
        # a vector made by another model is not comparable with this model's query vectors: treat it as missing
        r["emb"] = json.loads(r["emb"]) if r["emb"] and r["emb_model"] == current else None
        r["scope"] = json.loads(r["scope"]) if r["scope"] else None
        r["toks"] = tokens(r["text"])
    _cache[user] = (_version[user], rows)
    return rows


def bm25(query_toks: list[str], rows: list[dict], k1: float = 1.5, b: float = 0.75) -> list[float]:
    n = len(rows)
    df = Counter(t for r in rows for t in set(r["toks"]))
    avgdl = (sum(len(r["toks"]) for r in rows) / n) or 1.0
    scores = []
    for r in rows:
        tf = Counter(r["toks"])
        s = 0.0
        for t in set(query_toks):
            if t not in tf:
                continue
            idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
            s += idf * tf[t] * (k1 + 1) / (tf[t] + k1 * (1 - b + b * len(r["toks"]) / avgdl))
        scores.append(s)
    return scores


async def search(user: str, agent_id: str, query: str, k: int = 5, only: set[str] | None = None) -> list[dict]:
    """Top chunks for a question, from the documents this agent may see, with source and location."""
    # An expert agent searches only the documents it was built from; everyone else sees what they are allowed to.
    rows = [r for r in _load(user) if (r["doc_id"] in only if only is not None else (r["scope"] is None or agent_id in r["scope"]))]
    if not rows or not query.strip():
        return []
    kw = bm25(tokens(query), rows)
    qv = await memory.embed(query) if any(r["emb"] for r in rows) else None
    cos = [memory.cosine(qv, r["emb"]) if qv and r["emb"] else 0.0 for r in rows]

    def ranks(values: list[float], keep) -> dict[int, int]:
        order = sorted((i for i, v in enumerate(values) if keep(v)), key=lambda i: values[i], reverse=True)
        return {i: rank for rank, i in enumerate(order)}

    min_cos = embeddings.threshold("chunk")
    kw_rank, cos_rank = ranks(kw, lambda v: v > 0), ranks(cos, lambda v: v >= min_cos)
    fused: dict[int, float] = {}
    for rank_map in (kw_rank, cos_rank):  # reciprocal-rank fusion
        for i, rank in rank_map.items():
            fused[i] = fused.get(i, 0.0) + 1.0 / (60 + rank)
    top = sorted(fused, key=fused.get, reverse=True)[: max(k, rerank.CANDIDATES) if rerank.model_name() else k]
    scores = await rerank.rerank(query, [rows[i]["text"] for i in top]) if len(top) > 1 else None
    if scores:  # a cross-encoder reads question and passage together; keep its best k
        top = [i for _, i in sorted(zip(scores, top), key=lambda p: p[0], reverse=True)][:k]
    return [{"doc_id": rows[i]["doc_id"], "ord": rows[i].get("ord"), "doc": rows[i]["name"], "locator": rows[i]["locator"], "text": rows[i]["text"], "score": round(fused[i], 4), "cosine": round(cos[i], 3), "bm25": round(kw[i], 2)} for i in top]


def samples(user: str, doc_ids: list[str], per_doc: int = 3, chars: int = 350) -> list[dict]:
    """The first few passages of each document, used to draft an expert agent that fits the content."""
    out = []
    for d in list_docs(user):
        if d["id"] in doc_ids:
            rows = db.q("SELECT text, locator FROM kchunks WHERE doc_id=? ORDER BY ord LIMIT ?", (d["id"], per_doc))
            out.append({"name": d["name"], "kind": d["kind"], "passages": [r["text"][:chars] for r in rows]})
    return out


async def reindex(user: str | None = None) -> int:
    """Re-embed passages made with a different (or no) embedding model. Run at startup and after changing models."""
    current = embeddings.model_id()
    if current == "none":
        return 0
    where, args = "(c.emb IS NULL OR c.emb_model IS NULL OR c.emb_model != ?)", [current]
    if user:
        where += " AND d.user = ?"
        args.append(user)
    rows = db.q(f"SELECT c.id, c.text, d.user, d.id AS doc FROM kchunks c JOIN kdocs d ON d.id = c.doc_id WHERE {where}", tuple(args))
    done, users, docs = 0, set(), set()
    for i in range(0, len(rows), 32):
        batch = rows[i : i + 32]
        vecs = await embed_many([r["text"] for r in batch])
        for r, v in zip(batch, vecs):
            if v:
                db.run("UPDATE kchunks SET emb=?, emb_model=? WHERE id=?", (json.dumps(v), current, r["id"]))
                done += 1
                users.add(r["user"])
                docs.add(r["doc"])
    for d in docs:
        db.run("UPDATE kdocs SET embedded = (SELECT COUNT(*) FROM kchunks WHERE doc_id=? AND emb IS NOT NULL) WHERE id=?", (d, d))
    for u in users:
        _touch(u)
    return done


# ---------------- keeping the extracted text, so a document can be re-chunked later ----------------


def _save_segments(doc_id: str, segments: list[Segment]) -> None:
    with db._lock:
        conn = db.get()
        conn.execute("DELETE FROM ksegments WHERE doc_id=?", (doc_id,))
        conn.executemany("INSERT INTO ksegments (doc_id, ord, text, locator) VALUES (?,?,?,?)", [(doc_id, i, s.text, s.locator) for i, s in enumerate(segments)])
        conn.commit()


async def rechunk(user: str, doc_id: str, cfg: chunking.Config) -> dict:
    """Cut an existing document differently, without re-reading the file (OCR results are kept)."""
    rows = db.q("SELECT * FROM kdocs WHERE user=? AND id=?", (user, doc_id))
    if not rows:
        raise IngestError("No such document.")
    segs = [Segment(r["text"], r["locator"]) for r in db.q("SELECT text, locator FROM ksegments WHERE doc_id=? ORDER BY ord", (doc_id,))]
    if not segs:
        raise IngestError("This document was added before re-chunking existed. Upload it again to change its chunking.")
    try:
        chunks, used = await chunking.build_chunks(segs, cfg, rows[0]["kind"])
    except ValueError as e:
        raise IngestError(str(e)) from e
    chunks = chunks[:MAX_CHUNKS_PER_DOC]
    embs = await embed_many([c[0] for c in chunks])
    model = embeddings.model_id()
    with db._lock:
        conn = db.get()
        conn.execute("DELETE FROM kchunks WHERE doc_id=?", (doc_id,))
        conn.executemany("INSERT INTO kchunks (doc_id, ord, text, locator, emb, emb_model) VALUES (?,?,?,?,?,?)", [(doc_id, i, c[0], c[1], json.dumps(e) if e else None, model if e else None) for i, (c, e) in enumerate(zip(chunks, embs))])
        conn.execute("UPDATE kdocs SET chunks=?, embedded=?, chunking=? WHERE id=?", (len(chunks), sum(1 for e in embs if e), json.dumps(used), doc_id))
        conn.commit()
    _touch(user)
    return get_doc(user, doc_id) or {}


def preview(user: str, doc_id: str, limit: int = 8) -> list[dict]:
    """The first passages of a document, to see how it was cut."""
    if not db.q("SELECT 1 FROM kdocs WHERE user=? AND id=?", (user, doc_id)):
        return []
    return db.q("SELECT ord, locator, text FROM kchunks WHERE doc_id=? ORDER BY ord LIMIT ?", (doc_id, limit))
